"""Exporting a lab as a report you can hand over (0.9.0).

`lab.md` plus a folder of images pastes into most things, but a report goes to
someone as one file. Two formats:

    HTML   one self-contained page: the images are inside it (base64), so the
           file can be mailed, opened in any browser and printed to PDF
    Word   a .docx with the images embedded, headings as real Word headings,
           ready to paste into a client template or edit as it is

Both are built from the same reading of lab.json as lab.md (`model()`), so the
three never disagree about order: the findings summary first, then the report
top to bottom, each finding block under its section's heading followed by the
snips, notes and command output filed there.

The Word export needs python-docx (in requirements.txt). It is imported only
when used, so ShadowSnip runs without it and says what to install if it is
missing.
"""

from __future__ import annotations

import base64
import html
import io
from pathlib import Path

import lab

EXPORT_FORMATS = ("html", "docx")

# Severity colours, shared by both exports: background, text.
SEVERITY_COLOURS = {
    "Critical": ("#7a0019", "#ffffff"),
    "High": ("#d6332e", "#ffffff"),
    "Medium": ("#f0a020", "#1a1a1a"),
    "Low": ("#2f8cff", "#ffffff"),
    "Info": ("#6c6c78", "#ffffff"),
}

_IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}


# Said in place of a snip whose image is missing or not inside the lab folder.
MISSING = "image not found in the lab folder, left out"


class ExportError(RuntimeError):
    """An export that could not be written, with the reason."""


# -- the shared reading of the lab ---------------------------------------------------------------
def model(folder: Path) -> dict:
    """The report, in order, as plain data.

    {"name", "started", "findings": [...], "blocks": [...]}, where each block is
    one of:

        {"type": "heading", "level": 1-6, "text"}
        {"type": "finding", **finding}
        {"type": "snip", "label", "time", "caption", "path"}
        {"type": "note", "text", "evidence": [files]}
        {"type": "code", "label", "text"}
    """
    folder = Path(folder)
    state = lab._load_state(folder)
    if not state:
        raise ExportError(
            "This lab has no record yet. Take a snip or add a note first, with "
            "Keep a lab record switched on in Settings."
        )
    grouped = lab._ordered_groups(state)
    consumed = lab._attached_note_ids(grouped)
    stored = lab._findings_of(state)
    blocks: list[dict] = []

    def entries(here):
        for entry in here:
            if lab.entry_kind(entry) == "note":
                if str(entry.get("id", "")) in consumed:
                    continue
                blocks.append(_note_block(entry))
                continue
            filename = str(entry.get("file", ""))
            number = entry.get("number", 0)
            # Only an image that is really in the lab folder is read: see
            # lab.image_in. Anything else is reported in place, never embedded.
            image = lab.image_in(folder, filename)
            blocks.append({
                "type": "snip",
                "label": f"{number:03d}" if isinstance(number, int) else str(number),
                "time": str(entry.get("time", "")).strip(),
                "caption": str(entry.get("caption", "")).strip(),
                "path": str(image) if image is not None else "",
            })
            for note in here:
                if lab.entry_kind(note) == "note" and filename and filename in note.get("attach", ()):
                    block = _note_block(note)
                    block["evidence"] = []  # it sits right under its evidence
                    block["under_snip"] = True
                    blocks.append(block)

    entries(grouped.get("", []))
    for path, depth in lab._walk_sections(grouped):
        blocks.append({
            "type": "heading",
            "level": min(depth + 2, lab.MAX_HEADING_LEVEL),
            "text": path.rsplit(lab.SECTION_SEPARATOR, 1)[-1],
        })
        if path in stored:
            blocks.append({"type": "finding", **stored[path], "path": path})
        entries(grouped.get(path, []))

    return {
        "name": str(state.get("name", folder.name)),
        "started": str(state.get("started", "")).replace("T", " ")[:19],
        "findings": lab.findings_of_state(state),
        "blocks": blocks,
    }


def _note_block(entry: dict) -> dict:
    if lab.is_code(entry):
        return {"type": "code", "label": str(entry.get("label", "")).strip(),
                "text": str(entry.get("text", ""))}
    return {
        "type": "note",
        "text": str(entry.get("text", "")).strip(),
        "evidence": [str(n) for n in entry.get("attach", ()) if str(n).strip()],
    }


def default_path(folder: Path, fmt: str) -> Path:
    folder = Path(folder)
    return folder / f"{folder.name}-report.{fmt}"


def export(folder: Path, fmt: str, target: Path | None = None) -> Path:
    """Write the report for the lab in `folder`. Returns the file written."""
    if fmt not in EXPORT_FORMATS:
        raise ExportError(f"unknown export format '{fmt}'")
    data = model(folder)
    target = Path(target) if target else default_path(folder, fmt)
    payload = to_html(data) if fmt == "html" else to_docx(data)
    try:
        import storage

        storage.write_atomic(target, payload)
    except OSError as exc:
        raise ExportError(f"could not write {target}: {exc}") from exc
    return target


# -- HTML ---------------------------------------------------------------------------------------
_CSS = """
:root { --ink: #1d1d22; --dim: #5c5c66; --rule: #dcdce2; --code: #f4f4f7; --accent: #2f8cff; }
* { box-sizing: border-box; }
body { margin: 0; background: #ffffff; color: var(--ink);
  font: 15px/1.6 "Segoe UI", system-ui, -apple-system, "Helvetica Neue", Arial, sans-serif; }
main { max-width: 900px; margin: 0 auto; padding: 40px 32px 80px; }
h1 { font-size: 30px; margin: 0 0 4px; }
.started { color: var(--dim); margin: 0 0 28px; }
h2, h3, h4, h5, h6 { margin: 32px 0 12px; line-height: 1.3; }
h2 { font-size: 22px; border-bottom: 1px solid var(--rule); padding-bottom: 6px; }
h3 { font-size: 18px; } h4, h5, h6 { font-size: 16px; }
table.summary { border-collapse: collapse; width: 100%; margin: 8px 0 28px; }
.summary th, .summary td { text-align: left; padding: 7px 10px; border-bottom: 1px solid var(--rule); }
.summary th { color: var(--dim); font-weight: 600; font-size: 13px; }
.sev { display: inline-block; padding: 1px 9px; border-radius: 4px; font-size: 12px; font-weight: 600; }
.finding { border: 1px solid var(--rule); border-left: 5px solid var(--sev); border-radius: 6px;
  padding: 12px 16px; margin: 6px 0 18px; break-inside: avoid; }
.finding .head { display: flex; gap: 14px; align-items: center; margin-bottom: 6px; }
.finding .host { color: var(--dim); }
.finding h5 { margin: 12px 0 2px; font-size: 13px; text-transform: uppercase; letter-spacing: .04em; color: var(--dim); }
.finding p { margin: 0 0 4px; white-space: pre-wrap; }
figure { margin: 14px 0 6px; }
figure img { max-width: 100%; border: 1px solid var(--rule); border-radius: 4px; display: block; }
figcaption { color: var(--dim); font-size: 13px; margin-top: 4px; }
figcaption b { color: var(--ink); }
blockquote { margin: 6px 0 16px; padding: 4px 14px; border-left: 3px solid var(--accent); color: #33333b; white-space: pre-wrap; }
p.note { white-space: pre-wrap; }
.evidence { color: var(--dim); font-size: 13px; font-style: italic; }
.cmd { font: 600 13px/1.4 Consolas, "Cascadia Mono", "DejaVu Sans Mono", monospace; margin: 12px 0 4px; }
pre { background: var(--code); border: 1px solid var(--rule); border-radius: 6px; padding: 10px 12px;
  overflow-x: auto; font: 13px/1.45 Consolas, "Cascadia Mono", "DejaVu Sans Mono", monospace; margin: 4px 0 16px; }
.missing { color: #b00020; font-style: italic; }
@media print { main { padding: 0; } pre { white-space: pre-wrap; } }
"""


def _esc(text) -> str:
    return html.escape(str(text or ""), quote=True)


def _sev_chip(severity: str) -> str:
    bg, fg = SEVERITY_COLOURS.get(severity, ("#6c6c78", "#ffffff"))
    return f'<span class="sev" style="background:{bg};color:{fg}">{_esc(severity)}</span>'


def _data_uri(path: str) -> str | None:
    file = Path(path)
    try:
        data = file.read_bytes()
    except OSError:
        return None
    kind = _IMAGE_TYPES.get(file.suffix.lower(), "image/png")
    return f"data:{kind};base64,{base64.b64encode(data).decode('ascii')}"


def to_html(data: dict) -> bytes:
    out = [
        "<!doctype html>",
        '<html lang="en"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{_esc(data['name'])} report</title>",
        f"<style>{_CSS}</style></head><body><main>",
        f"<h1>{_esc(data['name'])}</h1>",
    ]
    if data["started"]:
        out.append(f'<p class="started">Started {_esc(data["started"])}</p>')
    if data["findings"]:
        out.append("<h2>Findings summary</h2>")
        out.append('<table class="summary"><thead><tr><th>#</th><th>Finding</th>'
                   "<th>Severity</th><th>Host</th></tr></thead><tbody>")
        for index, item in enumerate(data["findings"], start=1):
            out.append(
                f"<tr><td>{index}</td><td>{_esc(item['title'])}</td>"
                f"<td>{_sev_chip(item['severity'])}</td><td>{_esc(item['host'] or '-')}</td></tr>"
            )
        out.append("</tbody></table>")

    for block in data["blocks"]:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            out.append(f"<h{level}>{_esc(block['text'])}</h{level}>")
        elif kind == "finding":
            bg, _fg = SEVERITY_COLOURS.get(block["severity"], ("#6c6c78", "#fff"))
            out.append(f'<div class="finding" style="--sev:{bg}"><div class="head">'
                       f"<b>Finding</b> {_sev_chip(block['severity'])}")
            if block.get("host"):
                out.append(f'<span class="host">Affected host: {_esc(block["host"])}</span>')
            out.append("</div>")
            for field, heading in (("description", "Description"), ("impact", "Impact"),
                                   ("remediation", "Remediation")):
                if block.get(field):
                    out.append(f"<h5>{heading}</h5><p>{_esc(block[field])}</p>")
            out.append("</div>")
        elif kind == "snip":
            uri = _data_uri(block["path"]) if block["path"] else None
            label = f"<b>{_esc(block['label'])}</b>"
            bits = [label] + [_esc(x) for x in (block["time"], block["caption"]) if x]
            if uri is None:
                out.append(f'<p class="missing">Snip {_esc(block["label"])}: {MISSING}</p>')
            else:
                out.append(f'<figure><img src="{uri}" alt="Snip {_esc(block["label"])}">'
                           f"<figcaption>{' - '.join(bits)}</figcaption></figure>")
        elif kind == "note":
            if block.get("under_snip"):
                out.append(f"<blockquote>{_esc(block['text'])}</blockquote>")
            else:
                out.append(f'<p class="note">{_esc(block["text"])}</p>')
                if block["evidence"]:
                    out.append(f'<p class="evidence">Evidence: {_esc(", ".join(block["evidence"]))}</p>')
        elif kind == "code":
            if block["label"]:
                out.append(f'<div class="cmd">{_esc(block["label"])}</div>')
            out.append(f"<pre><code>{_esc(block['text'])}</code></pre>")
    out.append("</main></body></html>")
    return "\n".join(out).encode("utf-8")


# -- Word --------------------------------------------------------------------------------------
def docx_available() -> bool:
    import importlib.util

    return importlib.util.find_spec("docx") is not None


def _picture_stream(path: str):
    """The image as something Word can hold (PNG or JPEG), and its pixel size."""
    from PIL import Image

    with Image.open(path) as img:
        size = img.size
        if img.format in ("PNG", "JPEG"):
            return io.BytesIO(Path(path).read_bytes()), size
        buf = io.BytesIO()
        img.convert("RGBA" if "A" in img.getbands() else "RGB").save(buf, "PNG")
        buf.seek(0)
        return buf, size


def to_docx(data: dict) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise ExportError(
            "The Word export needs python-docx. Install it with:  "
            "pip install python-docx  (or pip install -r requirements.txt)"
        ) from exc

    def shade(cell_or_par, hex_fill: str) -> None:
        props = (cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc")
                 else cell_or_par._p.get_or_add_pPr())
        fill = OxmlElement("w:shd")
        fill.set(qn("w:val"), "clear")
        fill.set(qn("w:color"), "auto")
        fill.set(qn("w:fill"), hex_fill.lstrip("#"))
        props.append(fill)

    def mono(run) -> None:
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), "Consolas")

    doc = Document()
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    doc.add_heading(data["name"], level=0)
    if data["started"]:
        doc.add_paragraph(f"Started {data['started']}")

    if data["findings"]:
        doc.add_heading("Findings summary", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for cell, text in zip(table.rows[0].cells, ("#", "Finding", "Severity", "Host")):
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
        for index, item in enumerate(data["findings"], start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = item["title"]
            cells[3].text = item["host"] or "-"
            bg, fg = SEVERITY_COLOURS.get(item["severity"], ("#6c6c78", "#ffffff"))
            cells[2].text = ""
            run = cells[2].paragraphs[0].add_run(item["severity"])
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(fg.lstrip("#").upper())
            shade(cells[2], bg)
        doc.add_paragraph()

    for block in data["blocks"]:
        kind = block["type"]
        if kind == "heading":
            doc.add_heading(block["text"], level=max(1, min(block["level"] - 1, 9)))
        elif kind == "finding":
            bg, fg = SEVERITY_COLOURS.get(block["severity"], ("#6c6c78", "#ffffff"))
            table = doc.add_table(rows=2 if block.get("host") else 1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = ""
            table.rows[0].cells[0].paragraphs[0].add_run("Severity").bold = True
            table.rows[0].cells[1].text = ""
            run = table.rows[0].cells[1].paragraphs[0].add_run(block["severity"])
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(fg.lstrip("#").upper())
            shade(table.rows[0].cells[1], bg)
            if block.get("host"):
                table.rows[1].cells[0].text = ""
                table.rows[1].cells[0].paragraphs[0].add_run("Affected host").bold = True
                table.rows[1].cells[1].text = block["host"]
            doc.add_paragraph()
            for field, heading in (("description", "Description"), ("impact", "Impact"),
                                   ("remediation", "Remediation")):
                if block.get(field):
                    para = doc.add_paragraph()
                    para.add_run(heading).bold = True
                    for line in block[field].split("\n"):
                        doc.add_paragraph(line)
        elif kind == "snip":
            try:
                if not block["path"]:
                    raise FileNotFoundError(block["label"])
                stream, (w, h) = _picture_stream(block["path"])
            except Exception:  # noqa: BLE001 - a missing image is reported in place
                doc.add_paragraph().add_run(f"Snip {block['label']}: {MISSING}").italic = True
                continue
            # 96 dpi to inches, never wider than the page.
            width = min(Inches(w / 96), usable)
            doc.add_picture(stream, width=width)
            caption = doc.add_paragraph()
            caption.add_run(block["label"]).bold = True
            rest = " - ".join(x for x in (block["time"], block["caption"]) if x)
            if rest:
                caption.add_run(f" - {rest}")
            caption.runs[-1].font.size = Pt(9)
        elif kind == "note":
            for line in block["text"].split("\n"):
                para = doc.add_paragraph(line)
                if block.get("under_snip"):
                    para.paragraph_format.left_indent = Inches(0.3)
                    for run in para.runs:
                        run.italic = True
            if block.get("evidence"):
                para = doc.add_paragraph()
                run = para.add_run(f"Evidence: {', '.join(block['evidence'])}")
                run.italic = True
                run.font.size = Pt(9)
        elif kind == "code":
            if block["label"]:
                para = doc.add_paragraph()
                run = para.add_run(block["label"])
                run.bold = True
                mono(run)
            para = doc.add_paragraph()
            shade(para, "#f4f4f7")
            lines = block["text"].split("\n")
            for i, line in enumerate(lines):
                run = para.add_run(line)
                mono(run)
                if i < len(lines) - 1:
                    run.add_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
