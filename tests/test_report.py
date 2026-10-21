"""The pentest report (0.9.0): findings, command output, methodology, export."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog

import config
import lab
import report
import report_dialogs


@pytest.fixture(scope="module", autouse=True)
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def cfg_lab(cfg):
    lab.start(cfg, "box")
    cfg["active_lab"] = "box"
    return cfg


def _png(width=40, height=20, colour=(200, 30, 30)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buf, "PNG")
    return buf.getvalue()


def _md(cfg) -> str:
    return (lab.active_folder(cfg) / lab.INDEX_NAME).read_text(encoding="utf-8")


# -- findings ---------------------------------------------------------------------------------
def test_a_section_can_be_marked_as_a_finding(cfg_lab):
    lab.set_section(cfg_lab, "10.10.10.3/SMB")
    lab.save(_png(), "png", cfg_lab)
    stored = lab.set_finding(cfg_lab, "10.10.10.3 / SMB", {
        "severity": "high", "host": " 10.10.10.3 ", "description": "Anonymous share.",
        "impact": "Read and write.", "remediation": "Disable guest.",
    })
    assert stored["severity"] == "High" and stored["host"] == "10.10.10.3"
    assert lab.finding(cfg_lab, "10.10.10.3/SMB") == stored
    md = _md(cfg_lab)
    assert "## Findings summary" in md
    assert "| 1 | 10.10.10.3 / SMB | High | 10.10.10.3 |" in md
    block = md.index("> **Finding: High**")
    assert md.index("### SMB") < block < md.index("![001]")
    assert "**Remediation**\n\nDisable guest." in md


def test_a_finding_needs_a_known_severity_and_a_section(cfg_lab):
    assert lab.set_finding(cfg_lab, "Web", {"severity": "Spicy"}) is None
    assert lab.set_finding(cfg_lab, "", {"severity": "Low"}) is None
    assert lab.findings(cfg_lab) == []


def test_findings_sort_by_severity_then_report_order(cfg_lab):
    for path, sev in (("A", "Low"), ("B", "Critical"), ("C", "Low"), ("D", "Medium")):
        lab.set_finding(cfg_lab, path, {"severity": sev})
    assert [(f["path"], f["severity"]) for f in lab.findings(cfg_lab)] == [
        ("B", "Critical"), ("D", "Medium"), ("A", "Low"), ("C", "Low")]


def test_an_empty_finding_section_still_has_its_heading(cfg_lab):
    lab.set_finding(cfg_lab, "Web/Login", {"severity": "Info"})
    assert "### Login" in _md(cfg_lab)
    assert "Web/Login" in lab.sections(cfg_lab)


def test_removing_a_finding_keeps_the_section_and_evidence(cfg_lab):
    lab.set_section(cfg_lab, "SMB")
    lab.save(_png(), "png", cfg_lab)
    lab.set_finding(cfg_lab, "SMB", {"severity": "High"})
    assert lab.set_finding(cfg_lab, "SMB", None) is None
    md = _md(cfg_lab)
    assert "Finding" not in md and "## SMB" in md and "![001]" in md


def test_the_outline_carries_the_finding(cfg_lab):
    lab.set_finding(cfg_lab, "SMB", {"severity": "Critical", "host": "h"})
    sections = {s["path"]: s for s in lab.outline(cfg_lab)["sections"]}
    assert sections["SMB"]["finding"]["severity"] == "Critical"
    assert sections[""]["finding"] is None


# -- command output ---------------------------------------------------------------------------
def test_output_is_filed_as_a_code_block_with_its_command(cfg_lab):
    entry = lab.add_note(cfg_lab, "\r\n\n  Sharename  Type  \r\n  tmp        Disk\n\n",
                         fmt="code", label="smbclient  -N -L //10.10.10.3")
    assert entry["format"] == "code"
    assert entry["text"] == "  Sharename  Type\n  tmp        Disk"  # indentation kept
    assert entry["label"] == "smbclient -N -L //10.10.10.3"
    md = _md(cfg_lab)
    assert "`smbclient -N -L //10.10.10.3`\n\n```text\n  Sharename  Type\n  tmp        Disk\n```" in md


def test_the_fence_outgrows_backticks_in_the_output(cfg_lab):
    lab.add_note(cfg_lab, "echo ```nested```", fmt="code")
    assert "````text\necho ```nested```\n````" in _md(cfg_lab)


def test_output_attached_to_a_snip_renders_under_it(cfg_lab):
    path = lab.save(_png(), "png", cfg_lab)
    lab.add_note(cfg_lab, "uid=0(root)", attach=[path.name], fmt="code", label="id")
    md = _md(cfg_lab)
    assert md.index(f"![001]({path.name})") < md.index("`id`") < md.index("uid=0(root)")
    assert "> uid=0(root)" not in md


def test_editing_output_keeps_its_whitespace_and_emptying_removes_it(cfg_lab):
    entry = lab.add_note(cfg_lab, "a\n  b", fmt="code")
    assert lab.update_note(cfg_lab, entry["id"], "a\n    b\n") == "updated"
    assert lab.entries(cfg_lab)[-1]["text"] == "a\n    b"
    assert lab.update_note(cfg_lab, entry["id"], "  \n ") == "removed"


def test_the_outline_marks_output(cfg_lab):
    lab.add_note(cfg_lab, "out", fmt="code", label="cmd")
    node = lab.outline(cfg_lab)["sections"][0]["entries"][0]
    assert node["code"] is True and node["label"] == "cmd"


# -- methodology ------------------------------------------------------------------------------
def test_methodology_lays_out_the_phases_in_order(cfg_lab):
    added = lab.add_methodology(cfg_lab)
    assert added == list(lab.METHODOLOGY)
    assert lab.section(cfg_lab) == "Recon"
    md = _md(cfg_lab)
    positions = [md.index(f"## {phase}") for phase in lab.METHODOLOGY]
    assert positions == sorted(positions)
    assert lab.sections(cfg_lab)[:5] == list(lab.METHODOLOGY)
    assert lab.add_methodology(cfg_lab) == []  # once is enough


def test_phases_come_before_sections_used_earlier(cfg_lab):
    lab.set_section(cfg_lab, "Notes")
    lab.add_note(cfg_lab, "early")
    lab.add_methodology(cfg_lab)
    md = _md(cfg_lab)
    assert md.index("## Recon") < md.index("## Notes")
    assert lab.section(cfg_lab) == "Notes"  # a chosen section is not moved


def test_methodology_under_a_host(cfg_lab):
    added = lab.add_methodology(cfg_lab, under="10.10.10.3")
    assert added[0] == "10.10.10.3/Recon"
    assert "### Recon" in _md(cfg_lab)


def test_nothing_without_a_lab_record(cfg):
    assert lab.add_methodology(cfg) == []
    assert lab.set_finding(cfg, "A", {"severity": "Low"}) is None


# -- export -----------------------------------------------------------------------------------
@pytest.fixture
def full_lab(cfg_lab):
    lab.add_methodology(cfg_lab)
    lab.set_section(cfg_lab, "Enumeration/SMB")
    snip = lab.save(_png(), "png", cfg_lab)
    lab.set_caption(cfg_lab, snip.name, "Share list <b>")
    lab.add_note(cfg_lab, "Anonymous login.", attach=[snip.name])
    lab.add_note(cfg_lab, "  tmp  Disk", fmt="code", label="smbclient -L x")
    lab.set_finding(cfg_lab, "Enumeration/SMB", {"severity": "High", "host": "10.10.10.3",
                                                  "description": "Open share."})
    lab.set_section(cfg_lab, "Initial access")
    lab.save(_png(colour=(30, 200, 30)), "webp", cfg_lab)
    lab.set_finding(cfg_lab, "Initial access", {"severity": "Critical"})
    return cfg_lab


def test_the_model_follows_the_report_order(full_lab):
    data = report.model(lab.active_folder(full_lab))
    kinds = [(b["type"], b.get("text") or b.get("label") or b.get("severity")) for b in data["blocks"]]
    assert kinds[:4] == [("heading", "Recon"), ("heading", "Enumeration"), ("heading", "SMB"),
                         ("finding", "High")]
    assert any(b["type"] == "code" and b["label"] == "smbclient -L x" for b in data["blocks"])
    assert [f["severity"] for f in data["findings"]] == ["Critical", "High"]


def test_html_export_is_one_self_contained_file(full_lab):
    path = report.export(lab.active_folder(full_lab), "html")
    assert path.name == "box-report.html"
    text = path.read_text(encoding="utf-8")
    assert text.count("data:image/") == 2  # both snips inside the page
    assert "data:image/webp" in text
    assert "Share list &lt;b&gt;" in text  # escaped, never markup
    assert "Findings summary" in text and "Critical" in text
    assert "<pre><code>  tmp  Disk</code></pre>" in text


def test_docx_export_embeds_the_images(full_lab):
    pytest.importorskip("docx")
    import docx

    path = report.export(lab.active_folder(full_lab), "docx")
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media) == 2  # the webp was converted, not dropped
    document = docx.Document(str(path))
    texts = [p.text for p in document.paragraphs]
    assert "Findings summary" in texts and "smbclient -L x" in texts
    assert document.tables[0].rows[1].cells[2].text == "Critical"


def test_word_export_says_what_to_install_without_python_docx(full_lab, monkeypatch):
    import builtins

    real = builtins.__import__

    def no_docx(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ImportError(name)
        return real(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_docx)
    with pytest.raises(report.ExportError, match="pip install python-docx"):
        report.export(lab.active_folder(full_lab), "docx")


def test_a_folder_without_a_lab_record_cannot_export(tmp_path):
    with pytest.raises(report.ExportError):
        report.export(tmp_path, "html")


def test_a_new_lab_exports_an_empty_report(cfg_lab):
    text = report.export(lab.active_folder(cfg_lab), "html").read_text(encoding="utf-8")
    assert "<h1>box</h1>" in text and "Findings summary" not in text


# -- the dialogs ------------------------------------------------------------------------------
@pytest.mark.parametrize("text,command,output", [
    ("\u250c\u2500\u2500(jh\u327fkali)-[~]\n\u2514\u2500$ smbclient -N -L //x\nout", "smbclient -N -L //x", "out"),
    ("$ id\nuid=0", "id", "uid=0"),
    ("root@lame:~# whoami\nroot", "whoami", "root"),
    ("PS C:\\Users\\jh> whoami\njh", "whoami", "jh"),
    ("total 8\n-rw x", "", "total 8\n-rw x"),
])
def test_a_prompt_line_becomes_the_command(text, command, output):
    assert report_dialogs.split_command(text) == (command, output)


@pytest.mark.parametrize("path,host", [
    ("10.10.10.3/SMB", "10.10.10.3"), ("Web/app.example.com/Login", "app.example.com"),
    ("Enumeration/SMB", ""), ("10.10.10.3:8080", "10.10.10.3:8080"),
])
def test_the_host_is_guessed_from_the_breadcrumb(path, host):
    assert report_dialogs.guess_host(path) == host


def test_the_finding_dialog_fills_and_reads_back():
    dialog = report_dialogs.FindingDialog("10.10.10.3/SMB")
    assert dialog.host.text() == "10.10.10.3" and dialog.btn_remove is None
    dialog.severity.setCurrentText("Critical")
    dialog.fields["impact"].setPlainText("Root.")
    values = dialog.values()
    assert values["severity"] == "Critical" and values["impact"] == "Root."
    existing = report_dialogs.FindingDialog("SMB", {"severity": "Low", "host": "h"})
    assert existing.btn_remove is not None and existing.severity.currentText() == "Low"


def test_the_output_dialog_cannot_file_nothing():
    dialog = report_dialogs.OutputDialog("", "Recon", can_attach=False)
    assert not dialog.btn_file.isEnabled()
    dialog.output.setPlainText("x")
    assert dialog.btn_file.isEnabled()
    dialog.attach.setChecked(True)
    assert dialog.values() == ("", "x", False)  # nothing on screen to attach to


# -- through the app ----------------------------------------------------------------------------
@pytest.fixture
def shadow(tmp_path, monkeypatch):
    import app as app_mod
    import clipboard
    import platforms

    monkeypatch.setattr(config, "CONFIG_PATH", tmp_path / "cfg" / "config.json")
    monkeypatch.setattr(clipboard, "copy", lambda *a, **k: None)
    monkeypatch.setattr(platforms.HotkeyManager, "register", lambda *a, **k: None)
    s = app_mod.ShadowSnipApp(QApplication.instance())
    s.cfg["save_dir"] = str(tmp_path / "snips")
    s.cfg["lab_root"] = str(tmp_path / "labs")
    lab.start(s.cfg, "box")
    s.cfg["active_lab"] = "box"
    s._refresh_menu_text()
    yield s
    s.autocopy.release()
    s.preview.close()
    s.tray.hide()


def test_the_window_marks_the_current_section_as_a_finding(shadow, monkeypatch):
    lab.set_section(shadow.cfg, "10.10.10.3/SMB")

    def exec_(dialog):
        dialog.severity.setCurrentText("High")
        dialog.fields["description"].setPlainText("Open share")
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(report_dialogs.FindingDialog, "exec", exec_)
    shadow.edit_finding(None)
    stored = lab.finding(shadow.cfg, "10.10.10.3/SMB")
    assert stored["severity"] == "High" and stored["host"] == "10.10.10.3"

    monkeypatch.setattr(report_dialogs.FindingDialog, "exec", lambda d: report_dialogs.REMOVE)
    shadow.edit_finding("10.10.10.3/SMB")
    assert lab.finding(shadow.cfg, "10.10.10.3/SMB") is None


def test_the_clipboard_is_filed_as_output_after_a_look(shadow, monkeypatch):
    QApplication.clipboard().setText("$ whoami\nroot")
    seen = []

    def exec_(dialog):
        seen.append(dialog.values())
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(report_dialogs.OutputDialog, "exec", exec_)
    lab.set_section(shadow.cfg, "Post-exploitation")
    shadow.file_output()
    assert seen == [("whoami", "root", False)]
    entry = lab.entries(shadow.cfg)[-1]
    assert entry["format"] == "code" and entry["section"] == "Post-exploitation"


def test_cancelled_output_files_nothing(shadow, monkeypatch):
    monkeypatch.setattr(report_dialogs.OutputDialog, "exec", lambda d: QDialog.DialogCode.Rejected)
    shadow.file_output()
    assert lab.entries(shadow.cfg) == []


def test_export_from_the_window(shadow, monkeypatch, tmp_path):
    lab.add_note(shadow.cfg, "hello")
    opened = []
    import app as app_mod

    monkeypatch.setattr(app_mod.storage, "open_file", lambda p: opened.append(Path(p)))
    target = tmp_path / "out" / "report"  # no extension: the filter decides
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        lambda *a, **k: (str(target), "Word document (*.docx)"))
    shadow.export_report()
    assert opened == [target.with_suffix(".docx")]
    assert opened[0].is_file()


def test_new_labs_can_start_with_the_phases(shadow, monkeypatch):
    import app as app_mod

    shadow.cfg["lab_methodology"] = True
    monkeypatch.setattr(app_mod.QInputDialog, "getText", lambda *a, **k: ("fresh", True))
    shadow.start_lab()
    assert lab.sections(shadow.cfg)[:5] == list(lab.METHODOLOGY)


def test_the_output_hotkey_reaches_a_running_copy():
    import hotkey_linux
    import main

    assert hotkey_linux.ACTION_FLAGS["output"] == "--output"
    assert '"--output": b"output"' in hotkey_linux.TRIGGER_SOURCE
    assert config.DEFAULTS["output_hotkey"] == "ctrl+alt+o"
    assert ("output", "output_hotkey") in [(n, k) for n, k, _ in __import__("app").ShadowSnipApp.HOTKEYS]
    assert main.request_for(False, False, False) == b"show"


# -- a crafted lab cannot pull outside files into a report (0.9.1) ---------------------------------
def _crafted_lab(cfg, filenames):
    """A lab whose lab.json names `filenames`, as a lab shared by someone else could."""
    import json

    folder = lab.active_folder(cfg)
    state = lab._load_state(folder) or lab._blank_state(folder.name)
    for n, name in enumerate(filenames, start=1):
        state["entries"].append({"kind": "snip", "file": name, "number": n, "time": "", "section": ""})
    (folder / lab.STATE_NAME).write_text(json.dumps(state), encoding="utf-8")
    return folder


SECRET = b"\x89PNG\r\n\x1a\n" + b"TOP-SECRET-MARKER" * 20


def _leaks(html_text: str) -> bool:
    """Whether the secret file is in the page, as text or as an embedded image."""
    import base64

    return ("TOP-SECRET-MARKER" in html_text
            or base64.b64encode(SECRET).decode()[:40] in html_text
            or "data:image/" in html_text)


@pytest.mark.parametrize("make_name", [
    lambda outside: "../outside.png",
    lambda outside: "..\\outside.png",
    lambda outside: str(outside),            # absolute path
    lambda outside: "sub/../../outside.png",
])
def test_export_never_reads_a_file_outside_the_lab(cfg_lab, make_name):
    folder = lab.active_folder(cfg_lab)
    outside = folder.parent / "outside.png"
    outside.write_bytes(SECRET)
    _crafted_lab(cfg_lab, [make_name(outside)])
    html_text = report.export(folder, "html").read_text(encoding="utf-8")
    assert not _leaks(html_text)
    assert report.MISSING in html_text
    if report.docx_available():
        data = report.export(folder, "docx").read_bytes()
        assert b"TOP-SECRET-MARKER" not in data
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert not [n for n in z.namelist() if n.startswith("word/media/")]


def test_a_link_out_of_the_lab_is_refused(cfg_lab):
    import os

    folder = lab.active_folder(cfg_lab)
    outside = folder.parent / "outside.png"
    outside.write_bytes(SECRET)
    try:
        os.symlink(outside, folder / "001_link.png")
    except (OSError, NotImplementedError):
        pytest.skip("this system does not let the tests make a symlink")
    _crafted_lab(cfg_lab, ["001_link.png"])
    html_text = report.export(folder, "html").read_text(encoding="utf-8")
    assert not _leaks(html_text)
    assert report.MISSING in html_text


@pytest.mark.parametrize("name", ["", ".", "..", "../x.png", "a/b.png", "a\\b.png", "C:x.png",
                                  "notes.txt", "x.png\0.txt"])
def test_only_a_bare_image_name_in_the_folder_is_accepted(tmp_path, name):
    assert lab.image_in(tmp_path, name) is None


def test_a_real_snip_is_still_accepted(cfg_lab):
    path = lab.save(_png(), "png", cfg_lab)
    assert lab.image_in(lab.active_folder(cfg_lab), path.name) == path
    assert lab.image_in(lab.active_folder(cfg_lab), "999_missing.png") is None


def test_lab_md_never_links_outside_the_lab(cfg_lab):
    folder = _crafted_lab(cfg_lab, ["../outside.png", "C:\\secret.png", "002_ok.png"])
    md = lab.render_index(lab._load_state(folder))
    assert "../outside.png" not in md and "secret.png" not in md
    assert md.count("image left out") == 2
    assert "![003](002_ok.png)" in md
